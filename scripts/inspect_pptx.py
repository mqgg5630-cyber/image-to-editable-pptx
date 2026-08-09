#!/usr/bin/env python3
"""Inspect PPTX geometry and flag full-slide raster shortcuts."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def inspect(path: Path) -> dict:
    prs = Presentation(path)
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    slides = []
    total_oob = 0
    total_full_slide_images = 0
    for slide_index, slide in enumerate(prs.slides, start=1):
        text_shapes = 0
        picture_shapes = 0
        editable_chars = 0
        out_of_bounds = []
        full_slide_images = []
        for shape_index, shape in enumerate(slide.shapes, start=1):
            x, y, w, h = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
            if x < 0 or y < 0 or x + w > slide_w or y + h > slide_h:
                out_of_bounds.append(shape_index)
            if getattr(shape, "has_text_frame", False):
                text = shape.text or ""
                if text:
                    text_shapes += 1
                    editable_chars += len(text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_shapes += 1
                coverage = (w * h) / max(slide_w * slide_h, 1)
                if coverage >= 0.95:
                    full_slide_images.append(shape_index)
        total_oob += len(out_of_bounds)
        total_full_slide_images += len(full_slide_images)
        slides.append(
            {
                "slide": slide_index,
                "shape_count": len(slide.shapes),
                "text_shape_count": text_shapes,
                "editable_text_chars": editable_chars,
                "picture_count": picture_shapes,
                "out_of_bounds_shape_indices": out_of_bounds,
                "full_slide_picture_indices": full_slide_images,
            }
        )
    return {
        "file": str(path),
        "slide_count": len(prs.slides),
        "slide_width_emu": slide_w,
        "slide_height_emu": slide_h,
        "out_of_bounds_count": total_oob,
        "full_slide_picture_count": total_full_slide_images,
        "slides": slides,
        "pass": total_oob == 0 and total_full_slide_images == 0,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    report = inspect(Path(args.pptx))
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
